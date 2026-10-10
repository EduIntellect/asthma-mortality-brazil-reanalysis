#!/usr/bin/env python3
"""Rebuild Figure 2 (Kitagawa) and Supplementary Figure S1 (ages 5-34) from frozen
CSVs, and the author-review Word copy with all three figures embedded.
Run after build_docx.sh. No statistical model is fitted.
"""
from pathlib import Path
import csv

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
FIG = ROOT / "manuscript" / "figures"
BUILD = ROOT / "manuscript" / "article" / "build"

def csv_rows(path):
    with path.open(encoding="utf-8-sig", newline="") as fd:
        return list(csv.DictReader(fd))

def plot_kitagawa():
    data = csv_rows(ROOT / "data" / "kitagawa_endpoint_sensitivity.csv")
    assert len(data) == 6
    labels = [r["start_year"] + "\u2013" + r["end_year"] for r in data]
    comp = np.array([float(r["composition_effect"]) for r in data])
    rates = np.array([float(r["age_specific_rate_effect"]) for r in data])
    net = np.array([float(r["crude_rate_change"]) for r in data])
    assert np.allclose(comp + rates, net, atol=1e-11)
    fig, ax = plt.subplots(figsize=(7.1, 4.7))
    y = np.arange(6)
    ax.barh(y-0.16, comp, height=0.32, color="black", label="Age composition")
    ax.barh(y+0.16, rates, height=0.32, facecolor="white", edgecolor="black",
            hatch="///", label="Age-specific rates")
    ax.set_yticks(y, labels)
    ax.invert_yaxis()
    ax.set_xlim(-0.11, 0.225)
    ax.axvline(0, color="black", lw=0.8)
    ax.grid(axis="x", ls=":", lw=0.6, color="0.85")
    ax.set_axisbelow(True)
    ax.set_xlabel("Contribution to crude-rate change (deaths per 100,000)")
    ax.legend(frameon=False, loc="lower right")
    fig.tight_layout()
    output = FIG / "fig2_kitagawa_sensitivity.jpg"
    fig.savefig(output, dpi=300, pil_kwargs={"quality": 95})
    plt.close(fig)
    return output

def plot_young():
    data = csv_rows(ROOT / "primary_results.csv")
    a = sorted((r for r in data if r["series"] == "age_5_34"), key=lambda r:int(r["year"]))
    b = sorted((r for r in data if r["series"] == "age_7_34_control"), key=lambda r:int(r["year"]))
    assert len(a) == len(b) == 8
    years = np.array([int(r["year"]) for r in a])
    assert np.array_equal(years, [int(r["year"]) for r in b])
    rates = [float(r["rate"]) for r in a]
    control = [float(r["rate"]) for r in b]
    deaths = [int(r["n_deaths"]) for r in a]
    assert sum(deaths) == 1591
    fig, (top, bottom) = plt.subplots(2, 1, figsize=(7.1, 6.7), sharex=True,
                                     gridspec_kw={"height_ratios":[1.1, 0.85]})
    top.plot(years, rates, "k-o", lw=1.8, ms=5, label="Ages 5\u201334")
    top.plot(years, control, "k--s", lw=1.3, ms=4, mfc="white",
             label="Ages 7\u201334 (control)")
    top.set_ylabel("Deaths per 100,000")
    top.set_ylim(0.155, 0.28)
    top.legend(frameon=False, loc="upper left")
    top.grid(axis="y", ls=":", lw=0.6, color="0.84")
    bottom.bar(years, deaths, width=0.66, color="0.70", edgecolor="black")
    bottom.set_ylabel("Annual asthma deaths (ages 5\u201334)")
    bottom.set_ylim(0, max(deaths)*1.2)
    bottom.set_xticks(years)
    bottom.set_xlabel("Year")
    bottom.grid(axis="y", ls=":", lw=0.6, color="0.84")
    top.set_axisbelow(True)
    bottom.set_axisbelow(True)
    fig.tight_layout()
    output = FIG / "figS1_age_5_34_mortality.jpg"
    fig.savefig(output, dpi=300, pil_kwargs={"quality": 95})
    plt.close(fig)
    return output

def build_review_docx(p2, p3):
    base = BUILD / "manuscript.docx"
    result = BUILD / "manuscript_with_3_figures_REVIEW.docx"
    source = Document(base)
    assert len(source.inline_shapes) == 0, "Submission DOCX must have no embedded figures"
    captions = [
        ("Figure 1. Age-aligned national trends",
         "Relative change from 2014 for the crude and age-standardised national rates among completed ages 7 or more, alongside ages 5\u201334."),
        ("Figure 2. Kitagawa decomposition and endpoint sensitivity",
         "Post hoc age composition (black) and age-specific rates (hatched) across six pairs of years; components sum to the observed crude-rate change. Exact values in Supplementary Table S6."),
        ("Supplementary Figure S1. Mortality at ages 5\u201334 years",
         "Annual mortality rates for ages 5\u201334 and the overlapping ages 7\u201334 control; lower panel: annual deaths at ages 5\u201334."),
    ]
    plots = [FIG / "fig1_relative_change_ge7.jpg", p2, p3]
    assert all(p.is_file() for p in plots)
    source.add_page_break()
    title = source.add_paragraph("Figures \u2013 author review copy")
    title.runs[0].bold = True
    title.runs[0].font.size = Pt(14)
    source.add_paragraph("Figures are embedded here for internal review. The journal submission manuscript and its separately uploaded JPEG figures remain distinct.")
    for i, (path, (heading, caption)) in enumerate(zip(plots, captions)):
        if i:
            source.add_page_break()
        para = source.add_paragraph(heading)
        para.runs[0].bold = True
        para.runs[0].font.size = Pt(11)
        source.add_picture(str(path), width=Inches(6.0))
        info = source.add_paragraph(caption)
        for run in info.runs:
            run.font.size = Pt(9)
    source.save(result)
    check = Document(result)
    assert len(check.inline_shapes) == 3 and len(check.tables) == len(source.tables)
    return result

def main():
    FIG.mkdir(exist_ok=True, parents=True)
    BUILD.mkdir(exist_ok=True, parents=True)
    plt.rcParams.update({"font.family":"DejaVu Sans", "font.size":9,
                         "savefig.facecolor":"white", "figure.facecolor":"white"})
    p2 = plot_kitagawa()
    p3 = plot_young()
    result = build_review_docx(p2, p3)
    for p in [p2, p3, result]:
        print(p, p.stat().st_size, "bytes")

if __name__ == "__main__":
    main()
