"""Journal-style DOCX generation from reconciled Markdown and a marked comparison."""
from pathlib import Path
import re,difflib,json,hashlib,copy,unicodedata,sys
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
root=Path(__file__).resolve().parents[2];ed=root/'revision/editorial';out=root/'revision/deliverables'
def create():
 d=Document();s=d.sections[0];s.page_width=Inches(8.268);s.page_height=Inches(11.693);s.top_margin=s.bottom_margin=Inches(.8);s.left_margin=s.right_margin=Inches(.7);s.header_distance=s.footer_distance=Inches(.3)
 for name in ['Normal','Title','Heading 1','Heading 2','Heading 3','Caption']:
  st=d.styles[name];st.font.name='Liberation Serif';st.font.color.rgb=RGBColor(0,0,0)
  for el in st.element.xpath('.//w:pBdr'):el.getparent().remove(el)
  for fonts in st.element.xpath('.//w:rFonts'):
   for attr in list(fonts.attrib):
    if 'theme' in attr.lower():del fonts.attrib[attr]
 normal=d.styles['Normal'];normal.font.size=Pt(10.5);normal.paragraph_format.line_spacing=1.12;normal.paragraph_format.space_after=Pt(6)
 for name,size in [('Heading 1',15),('Heading 2',12),('Heading 3',11)]:
  st=d.styles[name];st.font.size=Pt(size);st.font.bold=True;st.paragraph_format.keep_with_next=True;st.paragraph_format.space_before=Pt(10);st.paragraph_format.space_after=Pt(5)
 d.styles['Title'].font.size=Pt(17);d.styles['Title'].font.bold=True
 d.styles['Caption'].font.size=Pt(9);d.styles['Caption'].font.italic=False
 footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=footer.add_run();field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field);fr=OxmlElement('w:r');ft=OxmlElement('w:t');ft.text='1';fr.append(ft);field.append(fr)
 d.core_properties.author='Francisco Federico García Crespi; Amaury de Souza';d.core_properties.subject='Prepared revision for author review';d.core_properties.comments='Generated reproducibly from revision/editorial Markdown and identified CSV outputs.'
 return d

def text_par(d,text,style=None):
 p=d.add_paragraph(style=style);p.add_run(text.replace('**',''));return p

def table(d,lines):
 cells=[[x.strip()for x in line.strip().strip('|').split('|')]for line in lines];cells=[cells[0]]+cells[2:];n=len(cells[0]);t=d.add_table(rows=1,cols=n);t.autofit=False
 # Fixed geometry and deliberate proportional widths, emphasising domain names and intervals.
 if cells[0][0]=='Endpoints':widths=[.9,.8,.8,.8,1.0,1.0,.9]
 elif cells[0][1]=='Family' and n==7:widths=[1.25,1.15,.65,1.25,.6,.5,.75]
 elif n==7:widths=[1.7,.52,.52,.57,.55,1.86,.55]
 elif n==10:widths=[.39,.84,.67,.49,.55,.57,.60,.60,.53,.53]
 elif n==8:widths=[1.15,1.2,.6,1.1,.5,.7,.85,.65]
 elif n==6:widths=[1.15,1.62,.52,.64,1.75,.54]
 elif n==4:widths=[1.0,1.4,2.0,1.65]
 else:widths=[6.85/n]*n
 total=sum(widths);widths=[w*6.85/total for w in widths]
 for c,w in zip(t.columns,widths):c.width=Inches(w)
 for j,v in enumerate(cells[0]):t.rows[0].cells[j].text=v
 trPr=t.rows[0]._tr.get_or_add_trPr();repeat=OxmlElement('w:tblHeader');trPr.append(repeat)
 for row in cells[1:]:
  c=t.add_row().cells
  for j,v in enumerate(row):c[j].text=v
 for i,row in enumerate(t.rows):
  pr=row._tr.get_or_add_trPr();ns=OxmlElement('w:cantSplit');pr.append(ns)
  for j,c in enumerate(row.cells):
   c.width=Inches(widths[j]);tcpr=c._tc.get_or_add_tcPr();margins=OxmlElement('w:tcMar')
   for side in ['top','bottom','left','right']:
    el=OxmlElement('w:'+side);el.set(qn('w:w'),'55');el.set(qn('w:type'),'dxa');margins.append(el)
   tcpr.append(margins);borders=OxmlElement('w:tcBorders')
   for side in ['top','bottom','left','right']:
    el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'3');el.set(qn('w:color'),'D6D6D6');borders.append(el)
   tcpr.append(borders)
   if i==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EEEEEE');tcpr.append(sh)
   for p in c.paragraphs:
    p.paragraph_format.space_after=Pt(1);p.paragraph_format.line_spacing=1.0
    for r in p.runs:r.font.size=Pt(8.3 if n>=7 else 9);r.bold=i==0
 d.add_paragraph().paragraph_format.space_after=Pt(1)
 return t

def read_md(path):
 d=create();lines=path.read_text().splitlines();i=0
 while i<len(lines):
  line=lines[i].strip()
  if not line:i+=1;continue
  if line.startswith('|'):
   block=[]
   while i<len(lines)and lines[i].strip().startswith('|'):block.append(lines[i].strip());i+=1
   table(d,block);continue
  image=re.match(r'!\[(.*?)\]\((.*?)\)',line)
  if image:
   p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run().add_picture(str(out/image[2]),width=Inches(6.7));text_par(d,image[1],'Caption');i+=1;continue
  if line.startswith('# '):text_par(d,line[2:],'Title')
  elif line.startswith('### '):text_par(d,line[4:],'Heading 2')
  elif line.startswith('## '):
   p=text_par(d,line[3:],'Heading 1')
   if re.match(r'## (Table 1\.|Figure [12]\.|Supplementary Figure|Supplementary methods)',line):p.paragraph_format.page_break_before=True
  elif line.startswith('- '):text_par(d,line[2:],'List Bullet')
  else:text_par(d,line)
  i+=1
 return d
build_list=[] if '--comparison-only' in sys.argv else [('Manuscript_CLEAN','manuscript.md'),('Supplement','supplement.md'),('Response_to_Amaury','response_to_amaury.md'),('Author_confirmations','author_confirmations.md')]
if '--documents'in sys.argv:
 selected=sys.argv[sys.argv.index('--documents')+1].split(',')
 assert set(selected)<=set(x[0]for x in build_list),'Unknown document'
 build_list=[x for x in build_list if x[0]in selected]
for name,src in build_list:
 d=read_md(ed/src);d.core_properties.title=name;d.save(out/f'{name}.docx')
# Complete textual comparison: deletions red/struck; insertions green; includes table-cell text.
def blocks(doc):
 texts=[]
 for b in doc.element.body:
  if b.tag.endswith('}p'):
   from docx.text.paragraph import Paragraph
   p=Paragraph(b,doc)
   if p.text.strip():texts.append(p.text)
  elif b.tag.endswith('}tbl'):
   from docx.table import Table
   for row in Table(b,doc).rows:texts.append(' | '.join(c.text for c in row.cells))
 return texts
reference=Document(root/'revision/evidence/reference_concordant.docx');old=blocks(reference);new=blocks(Document(out/'Manuscript_CLEAN.docx'))+blocks(Document(out/'Supplement.docx'))
d=create();text_par(d,'Comparison against the PDF sent to Amaury','Title');text_par(d,'Reference: 10 October 2026 PDF, SHA-256 fc552ab83d5254549fa4f1dfe880ab631a26c50bc8d861bfdc433331544e71fe. The concordant Word supplies reference text. This is a visibly marked textual comparison, not native Word Track Changes. Red strike-through = removed text; green = added text. Tables are represented as row text. Figures are regenerated from identified corrected results; the clean manuscript and supplement contain the final graphics.')
changes=[]
def compare_key(t):
 return ' '.join(unicodedata.normalize('NFKC',t).replace('“','\"').replace('”','\"').replace('’',"'").split())
for op,a,b,c,e in difflib.SequenceMatcher(None,[compare_key(t)for t in old],[compare_key(t)for t in new],autojunk=False).get_opcodes():
 if op=='equal':
  for t in new[c:e]:text_par(d,t)
 else:
  changes.append(dict(operation=op,reference_blocks=[a,b],revised_blocks=[c,e]))
  for t in old[a:b]:p=d.add_paragraph();r=p.add_run(t);r.font.color.rgb=RGBColor.from_string('9C2020');r.font.strike=True
  for t in new[c:e]:p=d.add_paragraph();r=p.add_run(t);r.font.color.rgb=RGBColor.from_string('15653A')
d.save(out/'Manuscript_MARKED_COMPARISON.docx');(ed/'comparison_blocks.json').write_text(json.dumps(changes,indent=2));(ed/'comparison.patch').write_text('\n'.join(difflib.unified_diff(old,new,fromfile='PDF-concordant Word reference',tofile='corrected manuscript and supplement',lineterm='')))
print(f'Created {len(build_list)} source documents and marked text comparison')
