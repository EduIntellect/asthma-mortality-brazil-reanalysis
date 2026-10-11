from pathlib import Path
import re,json,hashlib,unicodedata,subprocess
from docx import Document
from pypdf import PdfReader
import pandas as pd
root=Path(__file__).resolve().parents[2];ed=root/'revision/editorial';out=root/'revision/deliverables'
def norm(t):return re.sub(r'\s+','',unicodedata.normalize('NFKC',t).replace('“','"').replace('”','"').replace('’',"'").replace('−','-'))
def pdf_text(path):
 # Poppler sorts fallback-font glyphs by position (pypdf can append ∈ out of order).
 raw=subprocess.check_output(['pdftotext','-layout',str(path),'-'],text=True)
 return '\n'.join(re.sub(r'\n\s*\d+\s*$', '',page)for page in raw.split('\f'))
def source_blocks(path):
 lines=path.read_text().splitlines();parts=[];i=0
 while i<len(lines):
  s=lines[i].strip();i+=1
  if not s:continue
  if s.startswith('|'):
   group=[s]
   while i<len(lines)and lines[i].strip().startswith('|'):group.append(lines[i].strip());i+=1
   for line in [group[0]]+group[2:]:parts.extend(x.strip()for x in line.strip('|').split('|'))
  elif s.startswith('!['):parts.append(re.match(r'!\[(.*?)\]',s)[1])
  else:parts.append(re.sub(r'^#+ |^- ','',s).replace('**',''))
 return parts
report=[]
for stem,src,images in [('Manuscript_CLEAN','manuscript.md',2),('Supplement','supplement.md',1),('Response_to_Amaury','response_to_amaury.md',0),('Author_confirmations','author_confirmations.md',0)]:
 doc=Document(out/f'{stem}.docx');actual=[]
 for b in doc.element.body:
  if b.tag.endswith('}p'):
   from docx.text.paragraph import Paragraph
   text=Paragraph(b,doc).text
   if text.strip():actual.append(text)
  elif b.tag.endswith('}tbl'):
   from docx.table import Table
   for row in Table(b,doc).rows:actual.extend(c.text for c in row.cells)
 expected=source_blocks(ed/src);assert actual==expected,stem
 assert len(doc.inline_shapes)==images
 text=' '.join(actual);assert not re.search(r'\{\{|AUTHOR INPUT REQUIRED|AUTHOR CONFIRMATION REQUIRED|TODO|TBD',text)
 pdf=PdfReader(out/f'{stem}.pdf')
 # Content-stream order retains wrapped table-cell text; geometric order retains math glyphs.
 stream='\n'.join(re.sub(r'^\s*\d+\s*\n', '',p.extract_text()or'')for p in pdf.pages)
 flat=norm(pdf_text(out/f'{stem}.pdf')if stem=='Manuscript_CLEAN'else stream);pos=0;unmatched=[]
 for block in expected:
  if not norm(block):continue
  found=flat.find(norm(block),pos)
  if found<0:unmatched.append(block[:100])
  else:pos=found+len(norm(block))
 assert not unmatched,(stem,unmatched)
 report.append(dict(document=stem,pages=len(pdf.pages),markdown_to_docx_exact=True,pdf_text_order=True,figures=images,placeholders=False))
# Check every traced numerical value against its identified CSV row.
claims=json.loads((ed/'claim_trace.json').read_text())
for c in claims:
 df=pd.read_csv(root/c['source']).set_index('series');assert float(df.loc[c['series'],c['field']])==c['value']
# Reference Word/PDF ordered text check, including all table-cell paragraphs.
reference=Document(root/'revision/evidence/reference_concordant.docx');blocks=[]
for b in reference.element.body:
 if b.tag.endswith('}p'):
  from docx.text.paragraph import Paragraph
  text=Paragraph(b,reference).text
  if text.strip():blocks.append(text)
 elif b.tag.endswith('}tbl'):
  from docx.table import Table
  for row in Table(b,reference).rows:
   for cell in row.cells:
    blocks.extend(p.text for p in cell.paragraphs if p.text.strip())
flat=norm('\n'.join(re.sub(r'^\s*\d+\s*\n', '',p.extract_text()or'')for p in PdfReader(root/'revision/evidence/reference_sent.pdf').pages));pos=0
for b in blocks:
 index=flat.find(norm(b),pos);assert index>=0,b[:100];pos=index+len(norm(b))
assert len(blocks)==486
(root/'revision/qa/document_checks.json').write_text(json.dumps({'documents':report,'traced_values_checked':len(claims),'reference_blocks_verified':len(blocks)},indent=2));print('Document/source/PDF and reference concordance checks passed',len(claims),'numeric claims;',len(blocks),'reference blocks')
